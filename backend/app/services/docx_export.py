"""
Word (.docx) export built with python-docx.

Produces a valid, editable document: Title style for the note title, a
metadata block, Heading 1–4 for markdown headings, List Bullet / List Number
styles for lists, an indented italic quote style, and inline bold / italic /
monospace runs. Link targets are rendered as visible text rather than
executable markup. python-docx escapes all text nodes automatically, so note
content can never inject XML.
"""

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from app.content import parse_blocks, parse_inline

_MUTED = RGBColor(0x64, 0x74, 0x8F)
_CODE_FONT = "Consolas"


def _add_runs(paragraph, text):
    """Render inline markdown as runs (bold/italic/code/link-as-text)."""
    for kind, value in parse_inline(text):
        if kind == "strong":
            run = paragraph.add_run(value)
            run.bold = True
        elif kind == "em":
            run = paragraph.add_run(value)
            run.italic = True
        elif kind == "code":
            run = paragraph.add_run(value)
            run.font.name = _CODE_FONT
            run.font.size = Pt(9.5)
        elif kind == "link":
            label, url = value
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
            run.underline = True
        else:
            paragraph.add_run(value)


def _meta_paragraph(document, label, value):
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.size = Pt(9)
    label_run.font.color.rgb = _MUTED
    value_run = paragraph.add_run(value)
    value_run.font.size = Pt(9)
    value_run.font.color.rgb = _MUTED
    paragraph.paragraph_format.space_after = Pt(2)
    return paragraph


def build_docx(note):
    """Render a note to .docx bytes."""
    document = Document()

    # Base font for body text.
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    document.add_heading(note.title, level=0)

    category_name = note.category.name if note.category else "Uncategorized"
    tag_text = ", ".join(tag.name for tag in note.tags) or "—"

    def _fmt_dt(value):
        return value.strftime("%Y-%m-%d %H:%M UTC") if value else "—"

    _meta_paragraph(document, "Category", category_name)
    _meta_paragraph(document, "Tags", tag_text)
    _meta_paragraph(document, "Created", _fmt_dt(note.created_at))
    _meta_paragraph(document, "Last updated", _fmt_dt(note.updated_at))
    document.add_paragraph()

    blocks = parse_blocks(note.content)
    if not blocks:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("This note has no content.")
        run.italic = True
        return _pack(document)

    for block in blocks:
        if block["type"] == "heading":
            heading = document.add_heading("", level=min(block["level"], 4))
            _add_runs(heading, block["text"])
        elif block["type"] == "list":
            style = "List Number" if block["ordered"] else "List Bullet"
            for item in block["items"]:
                paragraph = document.add_paragraph(style=style)
                _add_runs(paragraph, item)
        elif block["type"] == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(14)
            run = paragraph.add_run(block["text"])
            run.italic = True
            run.font.color.rgb = _MUTED
        else:
            paragraph = document.add_paragraph()
            _add_runs(paragraph, block["text"])

    return _pack(document)


def _pack(document):
    """Serialize the document to bytes via an in-memory buffer."""
    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    buffer.close()
    return data
